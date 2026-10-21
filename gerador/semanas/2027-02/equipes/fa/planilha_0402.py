# -*- coding: utf-8 -*-
"""Planilha da pessoa no post de 04/02/2027, com as colunas do modelo que o próprio site baixa ("Baixar planilha
modelo", modelo-remessas.xlsx, guardado em explora/) e as 3 remessas em dólar da conta de exemplo (as mesmas da base:
12/07/2023, 14/03/2024 e 18/04/2024, US$ 4.000 cada; R$ 19.600, R$ 20.000 e R$ 20.400 debitados).
Uma planilha antiga, só com o total debitado: "Valor em R$" e "IOF" ficam em branco (são opcionais no modelo).

Sai em .csv (separador ";", números no formato brasileiro, UTF-8): a cópia local do site não traz o módulo de
planilhas .xlsx (/lib/xlsx.full.min.js, que o site carrega à parte), e o .csv o site lê com o leitor próprio.
O modal "Importar planilha" diz que aceita .xlsx, .xls e .csv.

python3 planilha_0402.py  ->  explora/minhas-remessas.csv
"""
import os
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
MODELO = os.path.join(AQUI, "explora", "modelo-remessas.xlsx")
SAIDA = os.path.join(AQUI, "explora", "minhas-remessas.csv")

CAB = ("Data", "Conta de origem", "Conta de destino", "Valor em moeda", "Valor em R$", "IOF", "TOTAL", "Moeda")
LINHAS = [  # na ordem do cabeçalho
    ("12/07/2023", "Conta corrente", "Corretora no exterior", "4.000,00", "", "", "19.600,00", "USD"),
    ("14/03/2024", "Conta corrente", "Corretora no exterior", "4.000,00", "", "", "20.000,00", "USD"),
    ("18/04/2024", "Conta corrente", "Corretora no exterior", "4.000,00", "", "", "20.400,00", "USD"),
]


def montar():
    ws = openpyxl.load_workbook(MODELO).worksheets[0]
    assert tuple(c.value for c in ws[1]) == CAB, [c.value for c in ws[1]]     # o mesmo cabeçalho do modelo do site
    txt = "\n".join(";".join(l) for l in [CAB] + LINHAS) + "\n"
    open(SAIDA, "w", encoding="utf-8-sig").write(txt)
    return SAIDA


if __name__ == "__main__":
    print(montar())
    print(open(SAIDA, encoding="utf-8-sig").read())
